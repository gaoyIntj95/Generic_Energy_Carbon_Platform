from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


INPUT = Path(sys.argv[1])
OUTPUT = Path(sys.argv[2])


def set_font(run, font_name: str = "SimHei", east_asia: str = "黑体") -> None:
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def set_heading_fonts(doc: Document) -> None:
    for level in range(1, 10):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "SimHei"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "SimHei")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "SimHei")
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")

    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            for run in paragraph.runs:
                set_font(run)


def next_num_id(numbering) -> int:
    values = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    return max(values, default=0) + 1


def make_number_instance(numbering, abstract_num_id: int, num_id: int) -> None:
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_num_id))
    num.append(abstract)
    numbering.append(num)


def assign_num(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = num_pr.find(qn("w:ilvl"))
    if ilvl is None:
        ilvl = OxmlElement("w:ilvl")
        num_pr.insert(0, ilvl)
    ilvl.set(qn("w:val"), "0")
    num = num_pr.find(qn("w:numId"))
    if num is None:
        num = OxmlElement("w:numId")
        num_pr.append(num)
    num.set(qn("w:val"), str(num_id))


def restart_process_lists(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        raise RuntimeError("No numbering definition found in source document")
    abstract_num_id = int(first_num.find(qn("w:abstractNumId")).get(qn("w:val")))
    num_id = next_num_id(numbering)
    process_count = 0
    active_num_id: int | None = None
    for paragraph in doc.paragraphs:
        if "Process 处理" in paragraph.text:
            active_num_id = num_id
            make_number_instance(numbering, abstract_num_id, active_num_id)
            num_id += 1
            process_count += 1
            continue
        if paragraph.style.name.startswith("Heading"):
            active_num_id = None
            continue
        if active_num_id is not None and paragraph.style.name == "List Number":
            assign_num(paragraph, active_num_id)
    return process_count


def main() -> None:
    doc = Document(str(INPUT))
    set_heading_fonts(doc)
    process_count = restart_process_lists(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(f"wrote {OUTPUT}; restarted {process_count} Process lists")


if __name__ == "__main__":
    main()
