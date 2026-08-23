from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


MD_PATH = Path(sys.argv[1])
INPUT_DOCX = Path(sys.argv[2])
OUTPUT_DOCX = Path(sys.argv[3])
ASSET_ROOT = MD_PATH.parent


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def parse_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        if not line.strip().startswith("|"):
            continue
        cells = [cell.strip().replace("\\|", "|") for cell in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            continue
        rows.append(cells)
    width = max((len(row) for row in rows), default=0)
    return [row + [""] * (width - len(row)) for row in rows]


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    for r_index, row in enumerate(rows):
        for c_index, value in enumerate(row):
            cell = table.cell(r_index, c_index)
            set_cell_margins(cell)
            if r_index == 0:
                set_cell_shading(cell, "EAF3F1")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.font.size = Pt(8.5 if len(value) > 50 else 9)
            if r_index == 0:
                run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_image(doc: Document, markdown_path: str, alt: str) -> None:
    image_path = (ASSET_ROOT / markdown_path).resolve()
    if not image_path.exists():
        raise FileNotFoundError(f"Missing image referenced by Markdown: {image_path}")
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.25))
    if alt:
        caption = doc.add_paragraph(alt)
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(6)
        for item in caption.runs:
            item.italic = True
            item.font.color.rgb = RGBColor(89, 89, 89)
            item.font.size = Pt(9)


def remove_body_after(doc: Document, start_paragraph_index: int) -> None:
    body = doc._body._element
    paragraphs = doc.paragraphs
    start_element = paragraphs[start_paragraph_index]._p
    removing = False
    for child in list(body):
        if child is start_element:
            removing = True
            continue
        if removing and child.tag != qn("w:sectPr"):
            body.remove(child)


def main() -> None:
    markdown = MD_PATH.read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(markdown) if line.startswith("## 1. Introduction"))
    source = markdown[start:]

    doc = Document(str(INPUT_DOCX))
    body_start = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "1. Introduction 简介")
    remove_body_after(doc, body_start)

    i = 0
    while i < len(source):
        line = source[i].rstrip()
        if not line.strip():
            i += 1
            continue

        heading = re.match(r"^(#{2,6})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1)) - 1
            paragraph = doc.add_paragraph(heading.group(2).strip(), style=f"Heading {level}")
            paragraph.paragraph_format.keep_with_next = True
            i += 1
            continue

        image = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line.strip())
        if image:
            add_image(doc, image.group(2), image.group(1))
            i += 1
            continue

        if line.strip().startswith("|"):
            table_lines = []
            while i < len(source) and source[i].strip().startswith("|"):
                table_lines.append(source[i].rstrip())
                i += 1
            add_table(doc, parse_table(table_lines))
            continue

        if line.strip() == "---":
            i += 1
            continue

        blockquote = re.match(r"^>\s*(.*)$", line)
        if blockquote:
            paragraph = doc.add_paragraph(blockquote.group(1).strip())
            paragraph.paragraph_format.left_indent = Inches(0.25)
            for run in paragraph.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(89, 89, 89)
            i += 1
            continue

        bullet = re.match(r"^\s*[-*]\s+(.*)$", line)
        number = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if bullet or number:
            style = "List Bullet" if bullet else "List Number"
            doc.add_paragraph((bullet or number).group(1).strip(), style=style)
            i += 1
            continue

        doc.add_paragraph(line.strip().replace("`", ""))
        i += 1

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT_DOCX))
    print(f"wrote {OUTPUT_DOCX}")


if __name__ == "__main__":
    main()
