from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from docx import Document


SRC = Path(r"C:\Users\340710\Desktop\通用工业企业能碳SaaS平台需求规格说明书_20260818.docx")
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "08-prd-20260818.md"
ASSET_DIR = ROOT / "docs" / "prd-assets"
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def clean(value: str) -> str:
    return " ".join(value.replace("\r", "").replace("\n", " ").split())


def md_cell(value: str) -> str:
    return clean(value).replace("|", "\\|")


def table_to_md(table) -> str:
    rows = [[md_cell(cell.text) for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n".join(lines)


def paragraph_images(paragraph, image_map):
    out = []
    for blip in paragraph._p.xpath('.//a:blip'):
        rid = blip.get('{%s}embed' % NS["r"])
        if rid in image_map:
            out.append(image_map[rid])
    return out


def main():
    if not SRC.exists():
        raise FileNotFoundError(SRC)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    image_map = {}
    with ZipFile(SRC) as zf:
        rels = ET.fromstring(zf.read("word/_rels/document.xml.rels"))
        for rel in rels:
            rid = rel.attrib.get("Id")
            target = rel.attrib.get("Target", "")
            if rid and target.startswith("media/"):
                name = Path(target).name
                dest = ASSET_DIR / name
                dest.write_bytes(zf.read("word/" + target))
                image_map[rid] = (Path("prd-assets") / name).as_posix()

    doc = Document(SRC)
    lines = [
        "# 通用工业企业能碳 SaaS 平台需求规格说明书",
        "",
        "> Markdown 备份来源：`通用工业企业能碳SaaS平台需求规格说明书_20260818.docx`。",
        "> 转换日期：2026-08-19。本文档按原始 DOCX 的正文顺序转换，保留标题、列表、表格及内嵌图片引用；原 Word 目录中的页码不作为 Markdown 内容复刻。",
        "> 说明：这是需求原文的 Markdown 化备份，不对业务规则作擅自改写；项目实现时仍应同时遵循 `docs/` 下的专项约束与验收文档。",
        "",
    ]
    body = doc.element.body
    table_by_el = {table._tbl: table for table in doc.tables}
    seen_tables = set()
    for child in body.iterchildren():
        if child.tag.endswith('}p'):
            paragraph = next((p for p in doc.paragraphs if p._p is child), None)
            if paragraph is None:
                continue
            text = clean(paragraph.text)
            imgs = paragraph_images(paragraph, image_map)
            style = paragraph.style.name if paragraph.style else "Normal"
            if text:
                if style.startswith("Heading"):
                    try:
                        level = int(style.split()[-1])
                    except ValueError:
                        level = 2
                    lines += ["#" * min(level + 1, 6) + " " + text, ""]
                elif style == "List Bullet":
                    lines.append("- " + text)
                else:
                    lines += [text, ""]
            for img in imgs:
                lines += [f"![文档内嵌图片]({img})", ""]
        elif child.tag.endswith('}tbl') and child in table_by_el:
            table = table_by_el[child]
            if id(table._tbl) not in seen_tables:
                rendered = table_to_md(table)
                if rendered:
                    lines += [rendered, ""]
                seen_tables.add(id(table._tbl))

    OUT.write_text("\n".join(lines).replace("\n\n\n", "\n\n"), encoding="utf-8")
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes), assets={len(image_map)}")


if __name__ == "__main__":
    main()
